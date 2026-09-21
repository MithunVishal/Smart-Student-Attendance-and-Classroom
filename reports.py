import io
import csv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def export_to_csv(records: list) -> io.StringIO:
    """Generates an in-memory CSV export of attendance records."""
    output = io.StringIO()
    writer = csv.writer(output)
    
    # Write header
    writer.writerow([
        'Register Number', 'Student Name', 'Department', 'Year', 'Section',
        'Date', 'IN Time', 'OUT Time', 'Duration', 'Status', 'Barcode ID'
    ])
    
    # Write data rows
    for r in records:
        writer.writerow([
            r.get('register_number', ''),
            r.get('student_name', ''),
            r.get('department', ''),
            r.get('year', ''),
            r.get('section', ''),
            r.get('date', ''),
            r.get('in_time', ''),
            r.get('out_time', '') or '-',
            r.get('duration_formatted', '') or '0s',
            r.get('status', ''),
            r.get('barcode_id', '')
        ])
        
    output.seek(0)
    return output

def export_to_excel(records: list, report_type: str = "Daily", date_label: str = "") -> io.BytesIO:
    """Generates a styled Microsoft Excel (.xlsx) workbook using openpyxl."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Records"
    
    # Fonts & Styles
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="4B5563")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10, color="1F2937")
    status_inside_font = Font(name="Calibri", size=10, bold=True, color="047857")
    status_exited_font = Font(name="Calibri", size=10, bold=True, color="4B5563")
    
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    alt_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    summary_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    # Title Block
    ws['A1'] = "Smart Student Attendance & Classroom Utilization System"
    ws['A1'].font = title_font
    
    ws['A2'] = f"Report Type: {report_type} Attendance Report | Period: {date_label or datetime.now().strftime('%Y-%m-%d')} | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws['A2'].font = subtitle_font
    
    # Summary Box
    total_records = len(records)
    inside_count = sum(1 for r in records if r.get('status') == 'Inside')
    exited_count = sum(1 for r in records if r.get('status') == 'Exited')
    unique_students = len(set(r.get('register_number') for r in records))
    
    ws['A4'] = "Total Sessions"
    ws['B4'] = total_records
    ws['C4'] = "Unique Students"
    ws['D4'] = unique_students
    ws['E4'] = "Currently Inside"
    ws['F4'] = inside_count
    ws['G4'] = "Exited"
    ws['H4'] = exited_count
    
    for col in range(1, 9):
        cell = ws.cell(row=4, column=col)
        cell.fill = summary_fill
        cell.font = Font(name="Calibri", size=10, bold=(col % 2 != 0), color="1E3A8A")
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if col % 2 == 0 else "left", vertical="center")
        
    # Table Headers
    headers = [
        "S.No", "Register No", "Student Name", "Department", "Year", "Sec",
        "Date", "IN Time", "OUT Time", "Duration", "Status", "Barcode ID"
    ]
    
    start_row = 6
    for col_idx, h_text in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    
    ws.row_dimensions[start_row].height = 26
    
    # Data Rows
    for idx, r in enumerate(records, 1):
        current_row = start_row + idx
        ws.row_dimensions[current_row].height = 20
        is_even = (idx % 2 == 0)
        
        row_values = [
            idx,
            r.get('register_number', ''),
            r.get('student_name', ''),
            r.get('department', ''),
            r.get('year', ''),
            r.get('section', ''),
            r.get('date', ''),
            r.get('in_time', ''),
            r.get('out_time', '') or '-',
            r.get('duration_formatted', '') or '0s',
            r.get('status', ''),
            r.get('barcode_id', '')
        ]
        
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            
            if is_even:
                cell.fill = alt_row_fill
                
            # Alignment & Status Coloring
            if col_idx in (1, 5, 6, 7, 8, 9, 10, 11, 12):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
            if col_idx == 11:  # Status column
                if val == 'Inside':
                    cell.font = status_inside_font
                else:
                    cell.font = status_exited_font
                    
    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or '')
            if cell.row < 4 and len(val_str) > 30:
                continue
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 11)
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def export_to_pdf(records: list, report_type: str = "Daily", date_label: str = "") -> io.BytesIO:
    """Generates a publication-grade PDF report using ReportLab."""
    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(letter),
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'RepTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1e40af'),
        alignment=1,
        fontName='Helvetica-Bold'
    )
    subtitle_style = ParagraphStyle(
        'RepSubTitle',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#4b5563'),
        alignment=1,
        fontName='Helvetica'
    )
    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#1f2937'),
        fontName='Helvetica'
    )
    cell_center = ParagraphStyle(
        'CellCenter',
        parent=cell_style,
        alignment=1
    )
    head_cell_style = ParagraphStyle(
        'HeadCell',
        parent=styles['Normal'],
        fontSize=9,
        leading=12,
        textColor=colors.white,
        alignment=1,
        fontName='Helvetica-Bold'
    )
    
    story = []
    
    # Title & Subtitle
    story.append(Paragraph("Smart Student Attendance & Classroom Utilization System", title_style))
    story.append(Spacer(1, 4))
    period_str = date_label or datetime.now().strftime('%Y-%m-%d')
    gen_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    story.append(Paragraph(f"Official {report_type} Attendance Audit Report &bull; Period: {period_str} &bull; Generated on: {gen_time}", subtitle_style))
    story.append(Spacer(1, 14))
    
    # Summary Box Table
    total_records = len(records)
    inside_count = sum(1 for r in records if r.get('status') == 'Inside')
    exited_count = sum(1 for r in records if r.get('status') == 'Exited')
    unique_students = len(set(r.get('register_number') for r in records))
    
    summary_data = [
        [
            Paragraph("<b>Total Sessions:</b> " + str(total_records), cell_center),
            Paragraph("<b>Unique Students:</b> " + str(unique_students), cell_center),
            Paragraph("<b>Currently Inside:</b> " + str(inside_count), cell_center),
            Paragraph("<b>Exited:</b> " + str(exited_count), cell_center)
        ]
    ]
    summary_table = Table(summary_data, colWidths=[180, 180, 180, 180])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#eff6ff')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#bfdbfe')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#dbeafe')),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 14))
    
    # Table Headers
    table_data = [[
        Paragraph("S.No", head_cell_style),
        Paragraph("Reg No", head_cell_style),
        Paragraph("Student Name", head_cell_style),
        Paragraph("Department", head_cell_style),
        Paragraph("Year / Sec", head_cell_style),
        Paragraph("Date", head_cell_style),
        Paragraph("IN Time", head_cell_style),
        Paragraph("OUT Time", head_cell_style),
        Paragraph("Duration", head_cell_style),
        Paragraph("Status", head_cell_style)
    ]]
    
    for idx, r in enumerate(records, 1):
        status_color = "#047857" if r.get('status') == 'Inside' else "#4b5563"
        status_html = f"<font color='{status_color}'><b>{r.get('status', '')}</b></font>"
        
        table_data.append([
            Paragraph(str(idx), cell_center),
            Paragraph(r.get('register_number', ''), cell_center),
            Paragraph(r.get('student_name', ''), cell_style),
            Paragraph(r.get('department', ''), cell_style),
            Paragraph(f"{r.get('year', '')} - {r.get('section', '')}", cell_center),
            Paragraph(r.get('date', ''), cell_center),
            Paragraph(r.get('in_time', ''), cell_center),
            Paragraph(r.get('out_time', '') or '-', cell_center),
            Paragraph(r.get('duration_formatted', '') or '0s', cell_center),
            Paragraph(status_html, cell_center)
        ])
        
    col_widths = [35, 75, 130, 125, 75, 70, 60, 60, 65, 55]
    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    
    t_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('TOPPADDING', (0, 0), (-1, 0), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
    ]
    
    # Alternating row background
    for i in range(1, len(table_data)):
        if i % 2 == 0:
            t_style.append(('BACKGROUND', (0, i), (-1, i), colors.HexColor('#f8fafc')))
            
    table.setStyle(TableStyle(t_style))
    story.append(table)
    
    doc.build(story)
    output.seek(0)
    return output
